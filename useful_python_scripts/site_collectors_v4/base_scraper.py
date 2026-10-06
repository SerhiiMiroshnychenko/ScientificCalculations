import requests
from bs4 import BeautifulSoup
from urllib.parse import urljoin, urlparse, urlunparse
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
import time
import argparse
import sys
from urllib3.util import Retry
from requests.adapters import HTTPAdapter

class BaseDocxScraper:
    def __init__(self, start_url, base_path_prefix=None, output_filename="output.docx", delay=0.8, retry=2, timeout=30, max_depth=50, split_by_segment_index=None):
        self.start_url = start_url
        parsed = urlparse(self.start_url)
        self.base_scheme = parsed.scheme or 'https'
        self.base_domain = parsed.netloc
        self.base_path_prefix = base_path_prefix
        self.output_filename = output_filename
        self.delay = delay
        self.retry = retry
        self.timeout = timeout
        self.max_depth = max_depth
        self.split_by_segment_index = split_by_segment_index

        self.visited = set()
        self.ordered_urls = []
        self.doc = Document()
        self.session = requests.Session()
        
        # Setup retries using HTTPAdapter
        retries = Retry(
            total=self.retry,
            backoff_factor=0.3,
            status_forcelist=[500, 502, 503, 504],
            raise_on_status=False
        )
        self.session.mount('http://', HTTPAdapter(max_retries=retries))
        self.session.mount('https://', HTTPAdapter(max_retries=retries))
        
        self.session.headers.update({
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
        })
        
        self._ensure_styles()

    def normalize_url(self, url, current_url=None):
        """
        Normalizes the URL:
        - Makes it absolute.
        - Strips anchors and query parameters.
        - Ignores static files.
        - Enforces domain boundaries.
        """
        if not url:
            return None

        low = url.strip().lower()
        if low.startswith(('mailto:', 'javascript:', 'tel:')):
            return None

        base = current_url or self.start_url
        absolute = urljoin(base, url)
        parsed = urlparse(absolute)
        cleaned = parsed._replace(fragment='', query='')

        if cleaned.netloc != self.base_domain:
            return None

        path = cleaned.path or '/'
        if not path.startswith('/'):
            path = '/' + path

        # Ignore static files
        if path.endswith(('.pdf', '.zip', '.tar', '.gz', '.md', '.png', '.jpg', '.jpeg', '.gif', '.svg')):
            return None

        # Check path prefix restriction if set
        if self.base_path_prefix:
            if not path.startswith(self.base_path_prefix):
                return None

        # Unify trailing slashes (strip them except for root path)
        if path != '/' and path.endswith('/'):
            path = path[:-1]

        normalized = urlunparse((self.base_scheme, self.base_domain, path, '', '', ''))
        return normalized

    def is_valid_url(self, url):
        """Checks if URL is internal and matches prefixes/rules."""
        if not url:
            return False
        parsed = urlparse(url)
        if parsed.scheme not in ('http', 'https'):
            return False
        if parsed.netloc != self.base_domain:
            return False
        if self.base_path_prefix and not parsed.path.startswith(self.base_path_prefix):
            return False
        if parsed.path.endswith(('.pdf', '.zip', '.tar', '.gz', '.md', '.png', '.jpg', '.jpeg', '.gif', '.svg')):
            return False
        return True

    def extract_links(self, soup, current_url):
        """Extracts internal links in order of appearance (first sidebar/nav, then content)."""
        links = []
        seen = set()

        def add_link(href):
            norm = self.normalize_url(href, current_url=current_url)
            if not norm or not self.is_valid_url(norm):
                return
            if norm in seen:
                return
            seen.add(norm)
            links.append(norm)

        # Look in sidebar/navigation menus first
        nav_menu = soup.find('nav') or soup.find('aside') or soup.find('div', class_=['sidebar', 'toctree', 'navigation'])
        if nav_menu:
            for a_tag in nav_menu.find_all('a', href=True):
                add_link(a_tag['href'])

        # Look in main content next
        main_content = (
            soup.find('main') or
            soup.find('article') or
            soup.find('div', class_='document') or
            soup.find('div', role='main') or
            soup.find('body')
        )
        if main_content:
            for a_tag in main_content.find_all('a', href=True):
                add_link(a_tag['href'])

        return links

    def _get_soup(self, url):
        """Fetches page content and parses it with BeautifulSoup."""
        resp = self.session.get(url, timeout=self.timeout)
        resp.raise_for_status()
        return BeautifulSoup(resp.content, 'html.parser')

    def extract_content(self, soup, url):
        """Extracts content, cleaning navigation panels, headers and footers."""
        content_data = {'title': '', 'text': [], 'url': url}

        # Remove header, footer, scripts, navigation elements
        for unwanted in soup.find_all(['nav', 'aside', 'header', 'footer', 'script', 'style']):
            unwanted.decompose()

        # Remove elements by typical navigation/ui classes
        for unwanted_class in [
            'sidebar', 'navigation', 'toctree', 'breadcrumb', 'menu',
            'navbar', 'footer', 'header', 'btn', 'button', 'prev-next', 'toc', 'pager',
            'search-results', 'search-box', 'social-media'
        ]:
            for element in soup.find_all(class_=lambda x: x and unwanted_class in str(x).lower()):
                element.decompose()

        main_content = (
            soup.find('main') or
            soup.find('article') or
            soup.find('div', class_='document') or
            soup.find('div', role='main') or
            soup.find('body')
        )

        if main_content:
            title_tag = main_content.find('h1')
            if title_tag:
                content_data['title'] = title_tag.get_text(strip=True)

            self._process_element(main_content, content_data['text'])

        return content_data

    def _process_element(self, element, text_list, parent_tag=None):
        """Recursively parses HTML elements to build structured document layout."""
        if not hasattr(element, 'name') or element.name is None:
            return

        # Skip administrative/navigational elements
        if element.name in ['nav', 'aside', 'header', 'footer', 'button', 'script', 'style']:
            return

        element_class = element.get('class')
        if element_class:
            classes = ' '.join(element_class).lower()
            skip_classes = [
                'sidebar', 'toctree', 'breadcrumb', 'navigation', 'menu',
                'navbar', 'btn', 'button', 'prev-next', 'footer', 'header', 'toc', 'pager'
            ]
            if any(skip in classes for skip in skip_classes):
                return

        # Paragraphs, Headings, Definition Lists
        if element.name in ['p', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'dt', 'dd']:
            # Skip paragraph blocks that contain only buttons/links styling
            links = element.find_all('a')
            if links:
                all_buttons = all(
                    (link.get('class') and any(btn in ' '.join(link.get('class')).lower() for btn in ['btn', 'button']))
                    for link in links if link.get('class')
                )
                if all_buttons and len(links) > 0:
                    return

            text = element.get_text(" ", strip=True)
            if text and len(text) > 2:
                text_list.append({'type': element.name, 'text': text})
            return

        # List items
        elif element.name == 'li':
            parent_ul = element.find_parent(['ul', 'ol'])
            if parent_ul:
                parent_class = parent_ul.get('class')
                if parent_class:
                    classes = ' '.join(parent_class).lower()
                    if any(skip in classes for skip in ['toctree', 'menu', 'nav', 'toc']):
                        return
            text = element.get_text(" ", strip=True)
            if text and len(text) > 2:
                text_list.append({'type': 'li', 'text': text})
            
            # Recurse for nested lists, code blocks or tables inside list items
            for child in element.children:
                if hasattr(child, 'name'):
                    self._process_element(child, text_list, element.name)
            return

        # Code Blocks (pre, code)
        elif element.name in ['pre', 'code']:
            if element.name == 'code' and element.find_parent('pre'):
                return

            detected_lang = None
            if element.name == 'pre':
                dl = element.get('data-language')
                if dl:
                    detected_lang = str(dl).strip().lower()

            if not detected_lang:
                classes_here = element.get('class') or []
                if classes_here:
                    classes_joined = ' '.join(classes_here).lower()
                    for prefix in ('language-', 'lang-'):
                        if prefix in classes_joined:
                            for token in classes_joined.split():
                                if token.startswith(prefix) and len(token) > len(prefix):
                                    detected_lang = token[len(prefix):]
                                    break
                            if detected_lang:
                                break

            # If inside Sphinx style container
            if not detected_lang:
                parent_hl = element.find_parent('div', class_=lambda x: x and 'highlight-' in ' '.join(x).lower())
                if parent_hl and parent_hl.get('class'):
                    cj = ' '.join(parent_hl.get('class')).lower()
                    for token in cj.split():
                        if token.startswith('highlight-') and len(token) > len('highlight-'):
                            detected_lang = token[len('highlight-'):]
                            break

            text = None
            if element.name == 'pre':
                # Expressive Code handling: <div class="ec-line"><div class="code">...</div></div>
                ec_lines = element.select('div.ec-line div.code')
                if ec_lines:
                    lines = []
                    for code_div in ec_lines:
                        line_text = code_div.get_text('', strip=False)
                        lines.append(line_text.rstrip())
                    text = '\n'.join(lines).strip('\n')
                else:
                    code_child = element.find('code')
                    if code_child:
                        text = code_child.get_text('', strip=False).strip('\n')

            if text is None:
                text = element.get_text('', strip=False).replace('\r', '')

            # Normalize whitespaces and lines
            lines = [ln.rstrip() for ln in text.splitlines()]
            while lines and not lines[0].strip():
                lines.pop(0)
            while lines and not lines[-1].strip():
                lines.pop()
            text = '\n'.join(lines)

            # Code Caption handling
            try:
                parent = element.parent
                if parent and getattr(parent, 'find', None):
                    cap_div = None
                    if 'literal-block-wrapper' in ' '.join((parent.get('class') or [])).lower():
                        cap_div = parent.find('div', class_=lambda x: x and 'code-block-caption' in ' '.join(x).lower())
                    if not cap_div:
                        cap_div = element.find_previous_sibling('div', class_=lambda x: x and 'code-block-caption' in ' '.join(x).lower())
                    if cap_div:
                        caption_text = cap_div.get_text(' ', strip=True)
                        if caption_text:
                            text_list.append({'type': 'code_caption', 'text': caption_text})
            except Exception:
                pass

            if text and len(text) > 2:
                text_list.append({'type': element.name, 'text': text, 'lang': detected_lang})
            return

        # Blockquotes
        elif element.name == 'blockquote':
            text = element.get_text(" ", strip=True)
            if text and len(text) > 10:
                text_list.append({'type': 'blockquote', 'text': text})
            return

        # Tables
        elif element.name == 'table':
            rows = []
            for tr in element.find_all('tr'):
                row = []
                for cell in tr.find_all(['th', 'td']):
                    row.append(cell.get_text(" ", strip=True))
                if any(row):
                    rows.append(row)
            if rows:
                text_list.append({'type': 'table', 'rows': rows})
            return

        # Generic structural containers
        elif element.name in ['div', 'section', 'article', 'main', 'body', 'figure', 'details', 'ul', 'ol', 'dl']:
            for child in element.children:
                if hasattr(child, 'name'):
                    self._process_element(child, text_list, element.name)
            return

    def add_to_document(self, content_data):
        """Appends scraped page structure to the docx document with styling rules."""
        if not content_data['text']:
            # Still append title and source URL for navigation index files
            p = self.doc.add_paragraph()
            run = p.add_run(f"Source: {content_data['url']}")
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(128, 128, 128)
            if content_data.get('title'):
                heading = self.doc.add_heading(content_data['title'], level=1)
                heading.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            self.doc.add_paragraph('_' * 80)
            self.doc.add_paragraph()
            return

        p = self.doc.add_paragraph()
        run = p.add_run(f"Source: {content_data['url']}")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(128, 128, 128)

        if content_data['title']:
            heading = self.doc.add_heading(content_data['title'], level=1)
            heading.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

        for item in content_data['text']:
            t = item.get('type')
            if t in ['h2', 'h3', 'h4', 'h5', 'h6']:
                level = int(t[1]) + 1
                self.doc.add_heading(item['text'], level=min(level, 9))
            
            elif t == 'code_caption':
                meta = self.doc.add_paragraph(item['text'])
                for r in meta.runs:
                    r.font.size = Pt(8)
                    r.font.color.rgb = RGBColor(120, 120, 120)
                    r.font.italic = True

            elif t in ['pre', 'code']:
                lang = item.get('lang')
                if lang:
                    meta = self.doc.add_paragraph(f"Code ({lang})")
                    for r in meta.runs:
                        r.font.size = Pt(8)
                        r.font.color.rgb = RGBColor(120, 120, 120)
                        r.font.italic = True

                p = self.doc.add_paragraph(item['text'])
                try:
                    p.style = self.doc.styles['Code Block']
                except KeyError:
                    p.style = 'No Spacing'
                    for r in p.runs:
                        r.font.name = 'Courier New'
                        r.font.size = Pt(9)
                        r.font.color.rgb = RGBColor(50, 50, 50)

            elif t == 'blockquote':
                p = self.doc.add_paragraph(item['text'])
                try:
                    p.style = self.doc.styles['Quote Block']
                except KeyError:
                    p.paragraph_format.left_indent = Pt(36)
                    p.paragraph_format.right_indent = Pt(36)
                    for r in p.runs:
                        r.font.italic = True

            elif t == 'li':
                self.doc.add_paragraph(item['text'], style='List Bullet')

            elif t == 'table':
                rows = item.get('rows') or []
                if rows:
                    cols = max(len(r) for r in rows)
                    if cols > 0:
                        table = self.doc.add_table(rows=0, cols=cols)
                        for style_name in ('Light List', 'Table Grid'):
                            try:
                                table.style = self.doc.styles[style_name]
                                break
                            except KeyError:
                                continue
                        for r in rows:
                            row_cells = table.add_row().cells
                            for i, cell_text in enumerate(r):
                                if i < len(row_cells):
                                    row_cells[i].text = cell_text
            else:
                txt = item.get('text', '')
                if txt and len(txt) > 5:
                    self.doc.add_paragraph(txt)

        self.doc.add_paragraph('_' * 80)
        self.doc.add_paragraph()

    def _collect_urls_recursive(self, url, depth=0):
        """Recursively scrapes structure mapping to collect unique URLs."""
        if url in self.visited or depth > self.max_depth:
            return
        
        self.visited.add(url)
        self.ordered_urls.append(url)

        if depth == 0:
            print(f"Scanning hierarchy structure [{len(self.ordered_urls)}]: {url}")

        try:
            soup = self._get_soup(url)
            links = self.extract_links(soup, url)

            for link in links:
                if link not in self.visited:
                    if depth < 2:
                        print(f"Scanning hierarchy structure [{len(self.ordered_urls)}]: {link}")
                    self._collect_urls_recursive(link, depth + 1)
            time.sleep(self.delay * 0.6)
        except Exception as e:
            print(f"Error scanning {url}: {str(e)}")

    def _get_grouped_filename(self, group_name):
        """Generates filename with group prefix/suffix."""
        base_parts = self.output_filename.rsplit('.', 1)
        if len(base_parts) == 2:
            return f"{base_parts[0]}_{group_name}.{base_parts[1]}"
        return f"{self.output_filename}_{group_name}"

    def scrape(self):
        """Triggers the full process: collecting links, scanning pages and saving (with splitting if configured)."""
        print(f"Starting scrape: {self.start_url}")
        print("Phase 1: Collecting list of URLs to download...")
        self._collect_urls_recursive(self.start_url)

        total_pages = len(self.ordered_urls)
        print(f"\nFound {total_pages} pages.")
        print("Phase 2: Downloading content from pages...\n")

        # Semantic splitting check
        if self.split_by_segment_index is not None:
            groups = {}
            for url in self.ordered_urls:
                parsed = urlparse(url)
                parts = parsed.path.strip('/').split('/')
                # parts looks like: ['documentation', '20.0', 'applications', 'finance', 'accounting']
                if len(parts) > self.split_by_segment_index:
                    group_name = parts[self.split_by_segment_index]
                else:
                    group_name = "general"
                
                # Clean group name (strip extension if it ends with .html)
                if group_name.endswith('.html'):
                    group_name = group_name.rsplit('.', 1)[0]
                
                if group_name not in groups:
                    groups[group_name] = []
                groups[group_name].append(url)

            # Process group by group
            for group_name, urls in groups.items():
                print(f"\n--- Processing group: {group_name} ({len(urls)} pages) ---")
                self.doc = Document()
                self._ensure_styles()
                
                for idx, url in enumerate(urls, 1):
                    try:
                        print(f"[{group_name}] Processing [{idx}/{len(urls)}]: {url}")
                        soup = self._get_soup(url)
                        content = self.extract_content(soup, url)
                        self.add_to_document(content)
                        time.sleep(self.delay)
                    except Exception as e:
                        print(f"Skipping {url} due to error: {str(e)}")
                        continue
                
                fn = self._get_grouped_filename(group_name)
                self.doc.save(fn)
                print(f"✓ Group {group_name} document saved: {fn}")
        else:
            # Traditional single file scrape
            for idx, url in enumerate(self.ordered_urls, 1):
                try:
                    print(f"Processing [{idx}/{total_pages}]: {url}")
                    soup = self._get_soup(url)
                    content = self.extract_content(soup, url)
                    self.add_to_document(content)
                    time.sleep(self.delay)
                except Exception as e:
                    print(f"Skipping {url} due to error: {str(e)}")
                    continue
            self.save()

        print(f"\nCompleted! Processed pages: {total_pages}")
        return total_pages

    def save(self):
        """Saves the output docx document (used only if splitting is disabled)."""
        if self.split_by_segment_index is None:
            self.doc.save(self.output_filename)
            print(f"✓ Document saved: {self.output_filename}")

    def _ensure_styles(self):
        """Sets custom Word styles for Code Block and Quote Block."""
        styles = self.doc.styles
        
        # Setup Code Block Style
        try:
            code_style = styles['Code Block']
        except KeyError:
            try:
                code_style = styles.add_style('Code Block', WD_STYLE_TYPE.PARAGRAPH)
            except Exception:
                code_style = None
        if code_style is not None:
            try:
                code_style.font.name = 'Courier New'
                code_style.font.size = Pt(9)
                code_style.font.color.rgb = RGBColor(50, 50, 50)
            except Exception:
                pass

        # Setup Quote Block Style
        try:
            quote_style = styles['Quote Block']
        except KeyError:
            try:
                quote_style = styles.add_style('Quote Block', WD_STYLE_TYPE.PARAGRAPH)
            except Exception:
                quote_style = None
        if quote_style is not None:
            try:
                quote_style.font.italic = True
                quote_style.font.color.rgb = RGBColor(80, 80, 80)
                pf = quote_style.paragraph_format
                pf.left_indent = Pt(36)
                pf.right_indent = Pt(36)
            except Exception:
                pass


def parse_args():
    parser = argparse.ArgumentParser(description="Universal Documentation Scraper to DOCX")
    parser.add_argument("--url", help="Start URL to scrape")
    parser.add_argument("--prefix", help="Path prefix validation")
    parser.add_argument("--output", help="Output filename")
    parser.add_argument("--delay", type=float, default=0.8, help="Request delay in seconds")
    parser.add_argument("--depth", type=int, default=50, help="Maximum recursive depth")
    parser.add_argument("--split-index", type=int, default=None, help="Segment index for semantic splitting")
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    if not args.url:
        print("Please provide a start URL via --url option. Use -h for help.")
        sys.exit(1)
        
    scraper = BaseDocxScraper(
        start_url=args.url,
        base_path_prefix=args.prefix,
        output_filename=args.output or "output.docx",
        delay=args.delay,
        max_depth=args.depth,
        split_by_segment_index=args.split_index
    )
    scraper.scrape()
